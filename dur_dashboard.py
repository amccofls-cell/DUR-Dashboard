from __future__ import annotations

import json, os, re, shutil, sqlite3, sys, threading, traceback
from pathlib import Path
from datetime import datetime
import tkinter as tk
from tkinter import ttk, filedialog, messagebox

from openpyxl import load_workbook

APP_NAME = "DUR Dashboard"
VERSION = "0.9.0"

CATEGORIES = [
    ("combo_paid", "병용금기 급여"),
    ("combo_unpaid", "병용금기 비급여"),
    ("age", "연령금기"),
    ("preg", "임부금기"),
    ("lact", "수유부주의"),
    ("dup", "효능군중복"),
    ("tele", "비대면진료 처방금지"),
    ("cost", "비용효과적인 함량"),
]
DUR_CARDS = [
    ("combo", "병용금기"), ("age", "연령금기"), ("preg", "임부금기"),
    ("lact", "수유부주의"), ("dup", "효능군중복"),
    ("tele", "비대면진료 금지"), ("cost", "비용효과적 함량"),
]

BG="#F5F7FB"; PANEL="#FFFFFF"; TEXT="#172033"; MUTED="#6B7280"; BORDER="#E5E7EB"
ACCENT="#4F46E5"; ACCENT_SOFT="#EEF2FF"; YES="#0F766E"; YES_BG="#ECFDF5"
NO="#64748B"; NO_BG="#F8FAFC"; WARN="#B45309"; WARN_BG="#FFF7ED"; ERROR="#B91C1C"


def app_dir() -> Path:
    root = Path(os.getenv("LOCALAPPDATA") or Path.home()) / "DUR_Dashboard"
    root.mkdir(parents=True, exist_ok=True)
    (root/"files").mkdir(exist_ok=True)
    return root

ROOT_DIR=app_dir(); CFG_PATH=ROOT_DIR/"config.json"; DB_PATH=ROOT_DIR/"dur_index.sqlite3"

def norm(v):
    if v is None: return ""
    s=str(v).strip().lower()
    s=re.sub(r"\s+", "", s)
    s=s.replace("㎖","ml").replace("밀리리터","ml").replace("밀리그램","mg")
    s=re.sub(r"[()\[\]{}·,._\-/]", "", s)
    return s

def sval(v):
    if v is None: return ""
    if isinstance(v, datetime): return v.strftime("%Y-%m-%d")
    return str(v).strip()

def unique_headers(values):
    out=[]; seen={}
    for i,v in enumerate(values):
        h=sval(v) or f"col_{i+1}"
        seen[h]=seen.get(h,0)+1
        if seen[h]>1: h=f"{h}_{seen[h]}"
        out.append(h)
    return out

def find_header_row(ws, max_scan=15):
    keys=("제품명","품목명","약품명","성분명","성분명a","저함량")
    best=(1,-1)
    for r in range(1,min(ws.max_row,max_scan)+1):
        vals=[norm(c.value) for c in ws[r][:min(ws.max_column,40)]]
        score=sum(any(k in x for k in keys) for x in vals)
        if score>best[1]: best=(r,score)
    return best[0]

def detect_name_cols(headers, category):
    hn=[norm(h) for h in headers]
    if category.startswith("combo"):
        a=[i for i,h in enumerate(hn) if h in ("제품명a","품목명a") or ("제품명" in h and h.endswith("a"))]
        b=[i for i,h in enumerate(hn) if h in ("제품명b","품목명b") or ("제품명" in h and h.endswith("b"))]
        return (a[:1],b[:1])
    cand=[]
    for i,h in enumerate(hn):
        if h in ("제품명","품목명","약품명") or h.startswith("제품명") or h.startswith("품목명") or h.startswith("약품명"):
            cand.append(i)
    return (cand[:2],[])

def cost_rows(ws):
    # 비용효과 파일은 보통 3~4행 다단 헤더. 실제 데이터 시작점을 찾고 저/고함량 그룹을 보존한다.
    rows=list(ws.iter_rows(min_row=1,max_row=min(ws.max_row,8),values_only=True))
    header_row=4 if len(rows)>=4 else 1
    for idx,row in enumerate(rows,1):
        vals=[norm(x) for x in row]
        if vals.count("제품코드")>=2 or ("제품코드" in vals and "제품명" in vals and "함량" in vals): header_row=idx
    group_row=max(1,header_row-1)
    groups=list(rows[group_row-1]) if len(rows)>=group_row else []
    base=list(rows[header_row-1])
    headers=[]; current=""
    for i,h in enumerate(base):
        g=sval(groups[i]) if i<len(groups) else ""
        if g: current=g
        prefix="저함량" if "저함량" in current else ("고함량" if "고함량" in current else "")
        name=sval(h) or f"col_{i+1}"
        headers.append((prefix+" "+name).strip())
    headers=unique_headers(headers)
    for row in ws.iter_rows(min_row=header_row+1, values_only=True):
        if not any(v is not None and sval(v)!="" for v in row): continue
        yield headers,row

def normal_rows(ws):
    hr=find_header_row(ws)
    headers=unique_headers([c.value for c in ws[hr]])
    for row in ws.iter_rows(min_row=hr+1,values_only=True):
        if not any(v is not None and sval(v)!="" for v in row): continue
        yield headers,row

class Indexer:
    def __init__(self, db_path=DB_PATH): self.db_path=db_path
    def conn(self):
        c=sqlite3.connect(self.db_path); c.execute("PRAGMA journal_mode=WAL"); return c
    def init(self):
        with self.conn() as c:
            c.execute("CREATE TABLE IF NOT EXISTS records(id INTEGER PRIMARY KEY, category TEXT, sheet TEXT, product TEXT, product_norm TEXT, side TEXT, data TEXT)")
            c.execute("CREATE INDEX IF NOT EXISTS ix_records_cat_name ON records(category,product_norm)")
            c.execute("CREATE TABLE IF NOT EXISTS status(category TEXT PRIMARY KEY, ok INTEGER, message TEXT, filename TEXT, indexed_at TEXT, sheet_count INTEGER, row_count INTEGER)")
    def clear_category(self,c,cat): c.execute("DELETE FROM records WHERE category=?",(cat,)); c.execute("DELETE FROM status WHERE category=?",(cat,))
    def index_file(self, category, path, progress=None):
        self.init(); path=Path(path); total=0; sheets=0
        try:
            wb=load_workbook(path,read_only=True,data_only=True)
            with self.conn() as c:
                self.clear_category(c,category)
                for ws in wb.worksheets:
                    sheets+=1
                    if progress: progress(f"{ws.title} 시트 읽는 중…")
                    iterator=cost_rows(ws) if category=="cost" else normal_rows(ws)
                    batch=[]
                    for headers,row in iterator:
                        d={headers[i]:sval(row[i]) if i<len(row) else "" for i in range(len(headers))}
                        if category=="cost":
                            # 저/고함량 제품명 각각 검색 가능하게 2개의 검색 엔트리 생성
                            low=next((d[k] for k in d if "저함량" in k and "제품명" in k),"")
                            high=next((d[k] for k in d if "고함량" in k and "제품명" in k),"")
                            for side,p in (("저함량",low),("고함량",high)):
                                if p: batch.append((category,ws.title,p,norm(p),side,json.dumps(d,ensure_ascii=False)))
                        else:
                            a,b=detect_name_cols(headers,category)
                            if category.startswith("combo"):
                                for side,idxs in (("A",a),("B",b)):
                                    for i in idxs:
                                        p=sval(row[i]) if i<len(row) else ""
                                        if p: batch.append((category,ws.title,p,norm(p),side,json.dumps(d,ensure_ascii=False)))
                            else:
                                for i in a:
                                    p=sval(row[i]) if i<len(row) else ""
                                    if p: batch.append((category,ws.title,p,norm(p),"",json.dumps(d,ensure_ascii=False)))
                        if len(batch)>=2000:
                            c.executemany("INSERT INTO records(category,sheet,product,product_norm,side,data) VALUES(?,?,?,?,?,?)",batch); total+=len(batch); batch=[]
                    if batch:
                        c.executemany("INSERT INTO records(category,sheet,product,product_norm,side,data) VALUES(?,?,?,?,?,?)",batch); total+=len(batch)
                c.execute("INSERT OR REPLACE INTO status VALUES(?,?,?,?,?,?,?)",(category,1,"정상",path.name,datetime.now().isoformat(timespec="seconds"),sheets,total))
            wb.close(); return True,f"{sheets}개 시트 · {total:,}개 검색 엔트리"
        except Exception as e:
            with self.conn() as c:
                self.clear_category(c,category)
                c.execute("INSERT OR REPLACE INTO status VALUES(?,?,?,?,?,?,?)",(category,0,str(e),path.name,datetime.now().isoformat(timespec="seconds"),sheets,total))
            return False,str(e)
    def statuses(self):
        self.init()
        with self.conn() as c: return {r[0]:r[1:] for r in c.execute("SELECT category,ok,message,filename,indexed_at,sheet_count,row_count FROM status")}
    def suggestions(self,q,limit=30):
        nq=norm(q)
        if not nq:return []
        with self.conn() as c:
            rows=c.execute("SELECT product, COUNT(*) n FROM records WHERE product_norm LIKE ? GROUP BY product ORDER BY CASE WHEN product_norm LIKE ? THEN 0 ELSE 1 END, length(product), product LIMIT ?",(f"%{nq}%",f"{nq}%",limit)).fetchall()
        return [x[0] for x in rows]
    def search_product(self,product):
        np=norm(product); out={k:[] for k,_ in CATEGORIES}
        with self.conn() as c:
            for cat,_ in CATEGORIES:
                rows=c.execute("SELECT sheet,product,side,data FROM records WHERE category=? AND product_norm=?",(cat,np)).fetchall()
                out[cat]=[{"sheet":r[0],"product":r[1],"side":r[2],"data":json.loads(r[3])} for r in rows]
        return out

class App(tk.Tk):
    def __init__(self):
        super().__init__(); self.title(f"{APP_NAME}  {VERSION}"); self.geometry("1360x820"); self.minsize(1120,700); self.configure(bg=BG)
        self.idx=Indexer(); self.idx.init(); self.cfg=self.load_cfg(); self.selected=None; self.after_id=None
        self.setup_style(); self.build(); self.refresh_status()
    def load_cfg(self):
        try:return json.loads(CFG_PATH.read_text(encoding="utf-8"))
        except:return {}
    def save_cfg(self): CFG_PATH.write_text(json.dumps(self.cfg,ensure_ascii=False,indent=2),encoding="utf-8")
    def setup_style(self):
        s=ttk.Style(self); s.theme_use("clam")
        s.configure("TFrame",background=BG); s.configure("Card.TFrame",background=PANEL)
        s.configure("TLabel",background=BG,foreground=TEXT,font=("Malgun Gothic",10))
        s.configure("Card.TLabel",background=PANEL,foreground=TEXT,font=("Malgun Gothic",10))
        s.configure("Title.TLabel",background=BG,foreground=TEXT,font=("Malgun Gothic",20,"bold"))
        s.configure("Sub.TLabel",background=BG,foreground=MUTED,font=("Malgun Gothic",9))
        s.configure("TButton",font=("Malgun Gothic",9),padding=(10,7))
        s.configure("Accent.TButton",background=ACCENT,foreground="white"); s.map("Accent.TButton",background=[("active","#4338CA")])
        s.configure("Treeview",font=("Malgun Gothic",9),rowheight=30,background=PANEL,fieldbackground=PANEL,borderwidth=0)
        s.configure("Treeview.Heading",font=("Malgun Gothic",9,"bold"),background="#F1F5F9",foreground=TEXT)
    def build(self):
        root=ttk.Frame(self); root.pack(fill="both",expand=True,padx=20,pady=18)
        top=ttk.Frame(root); top.pack(fill="x",pady=(0,14))
        ttk.Label(top,text="DUR Dashboard",style="Title.TLabel").pack(side="left")
        ttk.Label(top,text="원본 Excel 기반 · 로컬 전용",style="Sub.TLabel").pack(side="left",padx=14,pady=(9,0))
        self.global_status=ttk.Label(top,text="",style="Sub.TLabel"); self.global_status.pack(side="right",pady=(9,0))
        body=ttk.Frame(root); body.pack(fill="both",expand=True)
        self.left=ttk.Frame(body,style="Card.TFrame",width=300); self.left.pack(side="left",fill="y",padx=(0,14)); self.left.pack_propagate(False)
        self.main=ttk.Frame(body); self.main.pack(side="left",fill="both",expand=True)
        self.build_left(); self.build_main()
    def build_left(self):
        ttk.Label(self.left,text="기준파일 관리",style="Card.TLabel",font=("Malgun Gothic",12,"bold")).pack(anchor="w",padx=18,pady=(18,3))
        ttk.Label(self.left,text="새 파일을 등록하면 해당 항목만 교체됩니다.",style="Card.TLabel",foreground=MUTED,font=("Malgun Gothic",8)).pack(anchor="w",padx=18,pady=(0,12))
        self.slot_widgets={}
        for key,label in CATEGORIES:
            f=tk.Frame(self.left,bg=PANEL,highlightbackground=BORDER,highlightthickness=1); f.pack(fill="x",padx=14,pady=4)
            t=tk.Label(f,text=label,bg=PANEL,fg=TEXT,font=("Malgun Gothic",9,"bold"),anchor="w"); t.pack(fill="x",padx=10,pady=(8,0))
            st=tk.Label(f,text="미등록",bg=PANEL,fg=WARN,font=("Malgun Gothic",8),anchor="w",justify="left",wraplength=245); st.pack(fill="x",padx=10,pady=(2,5))
            b=ttk.Button(f,text="파일 등록/교체",command=lambda k=key:self.choose_file(k)); b.pack(anchor="e",padx=8,pady=(0,8))
            self.slot_widgets[key]=st
        ttk.Button(self.left,text="모든 파일 다시 인덱싱",command=self.reindex_all).pack(fill="x",padx=14,pady=14)
    def build_main(self):
        search=tk.Frame(self.main,bg=PANEL,highlightbackground=BORDER,highlightthickness=1); search.pack(fill="x")
        tk.Label(search,text="품목명 검색",bg=PANEL,fg=TEXT,font=("Malgun Gothic",11,"bold")).pack(anchor="w",padx=18,pady=(14,6))
        row=tk.Frame(search,bg=PANEL); row.pack(fill="x",padx=18,pady=(0,14))
        self.q=tk.StringVar(); e=tk.Entry(row,textvariable=self.q,font=("Malgun Gothic",13),relief="flat",bg="#F8FAFC",fg=TEXT,insertbackground=TEXT)
        e.pack(side="left",fill="x",expand=True,ipady=10); e.bind("<KeyRelease>",self.on_type); e.bind("<Return>",lambda _ : self.search_first())
        ttk.Button(row,text="검색",style="Accent.TButton",command=self.search_first).pack(side="left",padx=(10,0))
        self.suggest=tk.Listbox(self.main,height=5,font=("Malgun Gothic",10),relief="flat",highlightthickness=1,highlightbackground=BORDER,selectbackground=ACCENT,selectforeground="white")
        self.suggest.bind("<<ListboxSelect>>",self.pick_suggestion)
        # hidden until needed
        self.product_header=ttk.Frame(self.main); self.product_header.pack(fill="x",pady=(14,8))
        self.product_title=ttk.Label(self.product_header,text="품목을 검색해 주세요",font=("Malgun Gothic",15,"bold")); self.product_title.pack(side="left")
        self.product_sub=ttk.Label(self.product_header,text="",style="Sub.TLabel"); self.product_sub.pack(side="left",padx=12,pady=(5,0))
        cards=ttk.Frame(self.main); cards.pack(fill="x"); self.cards={}
        for i,(key,label) in enumerate(DUR_CARDS):
            f=tk.Frame(cards,bg=NO_BG,highlightbackground=BORDER,highlightthickness=1,width=150,height=84); f.grid(row=i//4,column=i%4,sticky="nsew",padx=(0 if i%4==0 else 7,0),pady=(0,7)); f.grid_propagate(False)
            tk.Label(f,text=label,bg=NO_BG,fg=MUTED,font=("Malgun Gothic",9)).pack(anchor="w",padx=12,pady=(11,3))
            val=tk.Label(f,text="—",bg=NO_BG,fg=NO,font=("Malgun Gothic",13,"bold")); val.pack(anchor="w",padx=12)
            f.bind("<Button-1>",lambda e,k=key:self.show_detail(k)); val.bind("<Button-1>",lambda e,k=key:self.show_detail(k))
            self.cards[key]=(f,val,label)
        for i in range(4): cards.columnconfigure(i,weight=1)
        detail=tk.Frame(self.main,bg=PANEL,highlightbackground=BORDER,highlightthickness=1); detail.pack(fill="both",expand=True,pady=(7,0))
        self.detail_title=tk.Label(detail,text="상세정보",bg=PANEL,fg=TEXT,font=("Malgun Gothic",11,"bold")); self.detail_title.pack(anchor="w",padx=16,pady=(12,6))
        cols=("항목","내용")
        self.tree=ttk.Treeview(detail,columns=cols,show="headings"); self.tree.heading("항목",text="항목"); self.tree.heading("내용",text="내용"); self.tree.column("항목",width=180,anchor="w"); self.tree.column("내용",width=760,anchor="w")
        sb=ttk.Scrollbar(detail,orient="vertical",command=self.tree.yview); self.tree.configure(yscrollcommand=sb.set); sb.pack(side="right",fill="y",pady=(0,10)); self.tree.pack(fill="both",expand=True,padx=(14,0),pady=(0,10))
    def refresh_status(self):
        sts=self.idx.statuses(); good=0
        for key,label in CATEGORIES:
            st=self.slot_widgets[key]; v=sts.get(key)
            if v and v[0]:
                good+=1; _,msg,fn,dt,sc,rc=v; st.config(text=f"✓ {fn}\n{sc}개 시트 · {rc:,}건",fg=YES)
            elif v: st.config(text=f"⚠ {v[2]}\n{v[1]}",fg=ERROR)
            else: st.config(text="미등록",fg=WARN)
        self.global_status.config(text=f"기준파일 {good}/8 정상")
    def choose_file(self,key):
        p=filedialog.askopenfilename(title="DUR Excel 선택",filetypes=[("Excel","*.xlsx")])
        if not p:return
        dest=ROOT_DIR/"files"/(key+"__"+Path(p).name)
        for old in (ROOT_DIR/"files").glob(key+"__*"):
            try:old.unlink()
            except:pass
        shutil.copy2(p,dest); self.cfg[key]=str(dest); self.save_cfg(); self.run_index(key,dest)
    def run_index(self,key,path):
        self.slot_widgets[key].config(text="인덱싱 중…",fg=ACCENT); self.update_idletasks()
        def worker():
            ok,msg=self.idx.index_file(key,path,lambda m:self.after(0,lambda:self.slot_widgets[key].config(text=m,fg=ACCENT)))
            self.after(0,lambda:self.index_done(key,ok,msg))
        threading.Thread(target=worker,daemon=True).start()
    def index_done(self,key,ok,msg): self.refresh_status(); messagebox.showinfo("인덱싱 완료" if ok else "인덱싱 오류",msg)
    def reindex_all(self):
        items=[(k,Path(v)) for k,v in self.cfg.items() if Path(v).exists()]
        if not items:return messagebox.showwarning("기준파일 없음","먼저 Excel 파일을 등록해 주세요.")
        def worker():
            for k,p in items: self.idx.index_file(k,p,lambda m,kk=k:self.after(0,lambda:self.slot_widgets[kk].config(text=m,fg=ACCENT)))
            self.after(0,lambda:(self.refresh_status(),messagebox.showinfo("완료","등록된 기준파일을 다시 인덱싱했습니다.")))
        threading.Thread(target=worker,daemon=True).start()
    def on_type(self,_=None):
        if self.after_id:self.after_cancel(self.after_id)
        self.after_id=self.after(180,self.update_suggestions)
    def update_suggestions(self):
        q=self.q.get().strip()
        if len(norm(q))<1:
            self.suggest.pack_forget(); return
        vals=self.idx.suggestions(q)
        self.suggest.delete(0,"end")
        for v in vals:self.suggest.insert("end",v)
        if vals:self.suggest.pack(fill="x",pady=(2,0),before=self.product_header)
        else:self.suggest.pack_forget()
    def pick_suggestion(self,_=None):
        sel=self.suggest.curselection()
        if not sel:return
        p=self.suggest.get(sel[0]); self.q.set(p); self.suggest.pack_forget(); self.do_search(p)
    def search_first(self):
        vals=self.idx.suggestions(self.q.get(),1)
        if vals:self.do_search(vals[0])
        else:messagebox.showinfo("검색 결과 없음","등록된 기준파일에서 일치하는 품목명을 찾지 못했습니다.\n파일 등록/인덱싱 상태를 확인해 주세요.")
    def do_search(self,p):
        self.selected=self.idx.search_product(p); self.selected_name=p; self.product_title.config(text=p)
        self.product_sub.config(text="등록된 DUR 원본 전체 통합 조회")
        sts=self.idx.statuses()
        mapping={"combo":["combo_paid","combo_unpaid"],"age":["age"],"preg":["preg"],"lact":["lact"],"dup":["dup"],"tele":["tele"],"cost":["cost"]}
        for key,label in DUR_CARDS:
            cats=mapping[key]; complete=all(sts.get(c) and sts[c][0] for c in cats); hits=sum(len(self.selected[c]) for c in cats)
            f,val,_=self.cards[key]
            if not complete: txt="확인불가"; bg=WARN_BG; fg=WARN
            elif hits:
                if key=="preg":
                    grades=sorted({self.find_value(x["data"],["금기등급","등급"]) for x in self.selected["preg"] if self.find_value(x["data"],["금기등급","등급"])})
                    txt="○ "+("/".join(g+"등급" for g in grades) if grades else "해당")
                elif key=="age":
                    x=self.selected["age"][0]["data"]; age=self.find_value(x,["특정연령"]); unit=self.find_value(x,["연령단위"]); cond=self.find_value(x,["연령처리조건"])
                    txt="○ "+(" ".join(z for z in [age+unit if age else "",cond] if z) or "해당")
                else: txt=f"○ 해당 ({hits})" if hits>1 else "○ 해당"
                bg=YES_BG; fg=YES
            else: txt="— 해당 없음"; bg=NO_BG; fg=NO
            f.config(bg=bg); val.config(text=txt,bg=bg,fg=fg)
            for w in f.winfo_children():
                if isinstance(w,tk.Label): w.config(bg=bg)
        first=next((k for k,_ in DUR_CARDS if self.card_has_hits(k)),"combo"); self.show_detail(first)
    def card_has_hits(self,key):
        m={"combo":["combo_paid","combo_unpaid"],"age":["age"],"preg":["preg"],"lact":["lact"],"dup":["dup"],"tele":["tele"],"cost":["cost"]}
        return any(self.selected.get(c) for c in m[key]) if self.selected else False
    def find_value(self,d,keys):
        for target in keys:
            for k,v in d.items():
                if norm(target)==norm(k) and v:return v
        return ""
    def show_detail(self,key):
        if not self.selected:return
        self.tree.delete(*self.tree.get_children()); label=dict(DUR_CARDS)[key]; self.detail_title.config(text=f"{label} 상세정보")
        cats={"combo":["combo_paid","combo_unpaid"],"age":["age"],"preg":["preg"],"lact":["lact"],"dup":["dup"],"tele":["tele"],"cost":["cost"]}[key]
        rows=[]
        for cat in cats:
            for hit in self.selected[cat]:
                d=hit["data"]; rows.append(("구분",dict(CATEGORIES)[cat])); rows.append(("원본 시트",hit["sheet"])); rows.append(("일치 품목",hit["product"]))
                if cat.startswith("combo"):
                    side=hit["side"]; other="B" if side=="A" else "A"
                    preferred=[f"성분명{side}",f"성분코드{side}",f"제품명{other}",f"품목명{other}",f"성분명{other}","고시번호","고시일자","상세정보"]
                    used=set()
                    for want in preferred:
                        for k,v in d.items():
                            if norm(k)==norm(want) and v and k not in used: rows.append((k,v)); used.add(k)
                    for k,v in d.items():
                        if v and k not in used: rows.append((k,v))
                else:
                    for k,v in d.items():
                        if v: rows.append((k,v))
                rows.append(("────────","────────"))
        if not rows: rows=[("결과","해당 없음")]
        for a,b in rows:self.tree.insert("","end",values=(a,b))

if __name__=="__main__":
    try: App().mainloop()
    except Exception:
        err=traceback.format_exc();
        try: messagebox.showerror("DUR Dashboard 오류",err)
        except: print(err)
